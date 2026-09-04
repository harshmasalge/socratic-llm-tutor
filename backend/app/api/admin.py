from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.security import OAuth2PasswordRequestForm
from sqlalchemy.ext.asyncio import AsyncSession
from ..dependencies import get_db
from ..core.security import verify_password, get_password_hash, create_access_token, get_current_admin
from ..schemas.admin import AdminLogin, Token, AppConfigUpdate, ConfigResponse
from ..models.domain import AppConfig
from ..models.domain import Student, Session, Message
from ..schemas.domain import StudentResponse, StudentDetailResponse, SessionResponse, MessageResponse
from sqlalchemy import select, desc, asc
from fastapi.responses import StreamingResponse
import csv
from datetime import datetime
from openpyxl import Workbook


router = APIRouter(tags=["admin"])

# Simple in‑memory admin credentials are stored in settings (see config.py)
# The login endpoint uses OAuth2PasswordRequestForm for compatibility with many clients

@router.post("/login", response_model=Token)
async def admin_login(form_data: OAuth2PasswordRequestForm = Depends(), db: AsyncSession = Depends(get_db)):
    # Settings holds hashed password; compare using utility
    from ..core.config import settings
    if form_data.username != settings.ADMIN_USERNAME:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Incorrect username")
    if not verify_password(form_data.password, settings.ADMIN_PASSWORD_HASH):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Incorrect password")
    access_token = create_access_token(data={"sub": settings.ADMIN_USERNAME})
    return Token(access_token=access_token, token_type="bearer")

# Protected endpoint to fetch current LLM configuration
@router.get("/config", response_model=ConfigResponse)
async def get_config(current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    # Load the two config rows using the AppConfig model
    from sqlalchemy import select
    result = await db.execute(
        select(AppConfig).where(AppConfig.config_key.in_(["LLM_MODEL", "SYSTEM_PROMPT"]))
    )
    rows = result.scalars().all()
    config_map = {row.config_key: row.config_value for row in rows}
    if "LLM_MODEL" not in config_map or "SYSTEM_PROMPT" not in config_map:
        raise HTTPException(status_code=404, detail="Config not found")
    return ConfigResponse(model=config_map["LLM_MODEL"], system_prompt=config_map["SYSTEM_PROMPT"]) 

# Protected endpoint to update configuration
@router.put("/config", response_model=ConfigResponse)
async def update_config(update: AppConfigUpdate, current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    # Upsert the two config entries using ORM
    from sqlalchemy import select
    # LLM model
    result = await db.execute(select(AppConfig).where(AppConfig.config_key == "LLM_MODEL"))
    model_entry = result.scalars().first()
    if model_entry:
        model_entry.config_value = update.model
    else:
        db.add(AppConfig(config_key="LLM_MODEL", config_value=update.model))
    # System prompt
    result = await db.execute(select(AppConfig).where(AppConfig.config_key == "SYSTEM_PROMPT"))
    prompt_entry = result.scalars().first()
    if prompt_entry:
        prompt_entry.config_value = update.system_prompt
    else:
        db.add(AppConfig(config_key="SYSTEM_PROMPT", config_value=update.system_prompt))
    await db.commit()
    return ConfigResponse(model=update.model, system_prompt=update.system_prompt)

# ----- New Admin Student Log Endpoints -----

@router.get("/students", response_model=list[StudentResponse])
async def admin_list_students(current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Student))
    return result.scalars().all()

@router.get("/students/{student_id}", response_model=StudentDetailResponse)
async def admin_get_student(student_id: str, current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Student).where(Student.id == student_id))
    student = result.scalars().first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
    sess_res = await db.execute(select(Session).where(Session.student_id == student_id).order_by(desc(Session.started_at)))
    sessions = sess_res.scalars().all()
    return StudentDetailResponse(**student.__dict__, sessions=sessions)

@router.get("/students/{student_id}/sessions", response_model=list[SessionResponse])
async def admin_get_student_sessions(student_id: str, current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Session).where(Session.student_id == student_id).order_by(desc(Session.started_at)))
    return result.scalars().all()

@router.get("/students/{student_id}/sessions/{session_id}", response_model=list[MessageResponse])
async def admin_get_session_messages(student_id: str, session_id: str, current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    sess_res = await db.execute(select(Session).where(Session.id == session_id, Session.student_id == student_id))
    if not sess_res.scalars().first():
        raise HTTPException(status_code=404, detail="Session not found for student")
    result = await db.execute(select(Message).where(Message.session_id == session_id).order_by(asc(Message.timestamp)))
    return result.scalars().all()

def _create_workbook(students, sessions, messages, include_summary=False):
    wb = Workbook()
    ws_students = wb.active
    ws_students.title = "Students"
    ws_students.append(["student_id", "name", "roll_no", "created_at"])
    for s in students:
        ws_students.append([s.id, s.name, s.roll_no, s.created_at])
    ws_sessions = wb.create_sheet(title="Sessions")
    ws_sessions.append(["session_id", "student_id", "started_at", "ended_at", "status", "message_count"])
    for sess in sessions:
        status = "ongoing" if not sess.ended_at else "completed"
        ws_sessions.append([sess.id, sess.student_id, sess.started_at, sess.ended_at, status, ""])
    ws_messages = wb.create_sheet(title="Messages")
    ws_messages.append(["message_id", "session_id", "student_id", "role", "content", "timestamp"])
    for m in messages:
        ws_messages.append([m.id, m.session_id, getattr(m.session, 'student_id', ''), m.role, m.content, m.timestamp])
    if include_summary:
        ws_summary = wb.create_sheet(title="Summary")
        ws_summary.append(["total_students", "total_sessions", "total_messages"])
        ws_summary.append([len(students), len(sessions), len(messages)])
    return wb

def _stream_workbook(wb: Workbook):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio

@router.get("/export/all")
async def export_all(
    format: str = "xlsx",
    start: str | None = None,
    end: str | None = None,
    current_admin: str = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db)
):
    # Validate format
    if format not in {"xlsx", "csv"}:
        raise HTTPException(status_code=400, detail="Invalid format")

    # Load all data
    stu_res = await db.execute(select(Student))
    sess_res = await db.execute(select(Session))
    students = stu_res.scalars().all()
    sessions = sess_res.scalars().all()

    # Base messages query with optional timestamp filter
    msg_query = select(Message).order_by(asc(Message.timestamp))
    if start:
        try:
            ts_start = datetime.fromisoformat(start)
            msg_query = msg_query.where(Message.timestamp >= ts_start)
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid start datetime")
    if end:
        try:
            ts_end = datetime.fromisoformat(end)
            msg_query = msg_query.where(Message.timestamp <= ts_end)
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid end datetime")
    msg_res = await db.execute(msg_query)
    messages = msg_res.scalars().all()

    # Filename construction
    date_range = "all"
    if start or end:
        start_part = start.split("T")[0] if start else ""
        end_part = end.split("T")[0] if end else ""
        date_range = f"{start_part}_{end_part}" if start and end else (f"{start_part}_" if start else f"_{end_part}")
    filename = f"logs_all_{date_range}.{format}"

    if format == "xlsx":
        wb = _create_workbook(students, sessions, messages, include_summary=True)
        bio = _stream_workbook(wb)
        return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename={filename}"})
    else:  # csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["message_id", "session_id", "student_id", "role", "content", "timestamp"])
        for m in messages:
            writer.writerow([m.id, m.session_id, getattr(m.session, 'student_id', ''), m.role, m.content, m.timestamp])
        bio = io.BytesIO(output.getvalue().encode())
        return StreamingResponse(bio, media_type="text/csv",
                                 headers={"Content-Disposition": f"attachment; filename={filename}"})

@router.get("/export/student/{student_id}")
async def export_student(student_id: str, current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    stu_res = await db.execute(select(Student).where(Student.id == student_id))
    student = stu_res.scalars().first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
    sess_res = await db.execute(select(Session).where(Session.student_id == student_id))
    sessions = sess_res.scalars().all()
    msg_res = await db.execute(select(Message).where(Message.session_id.in_([s.id for s in sessions])).order_by(asc(Message.timestamp)))
    messages = msg_res.scalars().all()
    wb = _create_workbook([student], sessions, messages, include_summary=False)
    bio = _stream_workbook(wb)
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=logs_student_{student_id}.xlsx"})

@router.get("/export/session/{session_id}")
async def export_session(session_id: str, current_admin: str = Depends(get_current_admin), db: AsyncSession = Depends(get_db)):
    sess_res = await db.execute(select(Session).where(Session.id == session_id))
    session_obj = sess_res.scalars().first()
    if not session_obj:
        raise HTTPException(status_code=404, detail="Session not found")
    msg_res = await db.execute(select(Message).where(Message.session_id == session_id).order_by(asc(Message.timestamp)))
    messages = msg_res.scalars().all()
    wb = _create_workbook([], [session_obj], messages, include_summary=False)
    bio = _stream_workbook(wb)
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=logs_session_{session_id}.xlsx"})

